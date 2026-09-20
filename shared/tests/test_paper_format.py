"""Regression tests for stable repository-wide DOCX layout invariants."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from docx import Document
from docx.document import Document as DocumentClass
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Twips


REPOSITORY_ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPOSITORY_ROOT / "shared"))

import paper_format  # noqa: E402
from paper_format import (apply_profile, load_profile, normalize_docx_file,
                          validate_docx_layout, validate_docx_math, validate_docx_tables)  # noqa: E402


def _pandoc_like_document():
    document = Document()
    title = document.add_paragraph("测试题名", style="Heading 1")
    major = document.add_paragraph("一、问题分析", style="Heading 2")
    secondary = document.add_paragraph("1.1 建模方法", style="Heading 3")
    tertiary = document.add_paragraph("1.1.1 数值方法", style="Heading 4")
    body = document.add_paragraph("正文")
    return document, title, major, secondary, tertiary, body


def test_pandoc_hierarchy_is_normalized_before_style_mapping() -> None:
    document, title, major, secondary, tertiary, _ = _pandoc_like_document()

    apply_profile(document, normalize_pandoc_headings=True)

    assert title.style.name == "Title"
    assert major.style.name == "Heading 1"
    assert secondary.style.name == "Heading 2"
    assert tertiary.style.name == "Heading 3"
    assert major.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert secondary.alignment == WD_ALIGN_PARAGRAPH.LEFT
    assert tertiary.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_keyword_character_indent_and_chapter_caption():
    document = Document()
    document.add_paragraph("测试题名")
    keywords = document.add_paragraph("关键词：储能；预测")
    caption = document.add_paragraph("表3-1 分章结果", style="Caption")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "结果"
    apply_profile(document)
    assert keywords._p.pPr.ind.get(qn("w:firstLineChars")) == "0"
    assert caption.text == "表3-1　分章结果"

def test_appendix_letter_caption_is_normalized_and_not_a_duplicate():
    document = Document()
    document.add_paragraph("测试题名")
    chapter = document.add_paragraph("表1-1 分章结果", style="Caption")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "结果"
    appendix = document.add_paragraph("表A-1 附录清单", style="Caption")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "内容"
    apply_profile(document)
    assert chapter.text == "表1-1　分章结果"
    assert appendix.text == "表A-1　附录清单"
    assert paper_format.validate_docx_captions(document) == []


def test_non_pandoc_h1_h2_hierarchy_is_not_rewritten() -> None:
    document = Document()
    major = document.add_paragraph("一、实际一级标题", style="Heading 1")
    secondary = document.add_paragraph("1.1 实际二级标题", style="Heading 2")

    apply_profile(document)

    assert major.style.name == "Heading 1"
    assert secondary.style.name == "Heading 2"
    assert major.alignment == WD_ALIGN_PARAGRAPH.CENTER
    assert secondary.alignment == WD_ALIGN_PARAGRAPH.LEFT


def test_caption_separator_preserves_split_runs_and_body_character_indent() -> None:
    document = Document()
    body = document.add_paragraph("普通正文")
    caption = document.add_paragraph()
    for text in ("表", "12", " ", "English 标题"):
        caption.add_run(text)
    table = document.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "数值 1.35"
    profile = load_profile()
    apply_profile(document, profile)
    assert caption.text == "表12" + profile["typography"]["caption"]["label_separator"] + "English 标题"
    assert caption.style.name == "Table Caption"
    assert all(run.bold for run in caption.runs)
    assert caption.runs[0].text == "表"
    assert caption.runs[1].text == "12"  # never move a field's cached number into its label
    assert body._p.pPr.ind.get(qn("w:firstLineChars")) == "200"
    assert table.cell(0, 0).paragraphs[0].style.name == "Table Text"
    assert table.cell(0, 0).paragraphs[0]._p.pPr.ind.get(qn("w:firstLineChars")) == "0"
    assert document.styles["Heading 2"].element.pPr.ind.get(qn("w:firstLineChars")) == "0"


def test_layout_validator_enforces_portrait_headings_and_body_spacing() -> None:
    document, _, level_2, level_3, level_4, body = _pandoc_like_document()
    profile = load_profile()
    apply_profile(document, profile)

    assert validate_docx_layout(document, profile) == []

    level_2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    level_3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    level_4.alignment = WD_ALIGN_PARAGRAPH.CENTER
    body.paragraph_format.line_spacing = 1.0
    document.sections[0].orientation = WD_ORIENT.LANDSCAPE
    document.sections[0].page_width, document.sections[0].page_height = (
        document.sections[0].page_height,
        document.sections[0].page_width,
    )

    errors = validate_docx_layout(document, profile)
    assert any("Heading 2" in error for error in errors)
    assert any("Heading 3" in error for error in errors)
    assert any("Heading 4" in error for error in errors)
    assert any("正文行距" in error for error in errors)
    assert any("page_width" in error or "page_height" in error for error in errors)


def test_layout_validator_rejects_landscape_flag_with_portrait_dimensions() -> None:
    document = Document()
    document.add_paragraph("正文")
    profile = load_profile()
    apply_profile(document, profile)
    section = document.sections[0]
    portrait_dimensions = (section.page_width.twips, section.page_height.twips)

    section.orientation = WD_ORIENT.LANDSCAPE

    assert (section.page_width.twips, section.page_height.twips) == portrait_dimensions
    errors = validate_docx_layout(document, profile)
    assert any("orientation" in error for error in errors)

    apply_profile(document, profile)
    assert section.orientation == WD_ORIENT.PORTRAIT
    assert validate_docx_layout(document, profile) == []


def _row_border_exception(row):
    exception = OxmlElement("w:tblPrEx")
    borders = OxmlElement("w:tblBorders")
    edge = OxmlElement("w:bottom")
    edge.set(qn("w:val"), "single")
    borders.append(edge)
    exception.append(borders)
    width = OxmlElement("w:tblW")
    width.set(qn("w:type"), "auto")
    exception.insert(0, width)
    row._tr.insert(0, exception)
    return exception


def test_row_border_exception_is_removed_but_other_row_layout_is_preserved():
    document = Document()
    table = document.add_table(rows=3, cols=2)
    exception = _row_border_exception(table.rows[1])
    apply_profile(document)
    assert exception.find(qn("w:tblBorders")) is None
    assert exception.find(qn("w:tblW")) is not None
    assert validate_docx_tables(document) == []

    profile = load_profile()
    profile["tables"]["remove_row_tblPrEx_borders"] = False
    retained = _row_border_exception(table.rows[2])
    apply_profile(document, profile)
    assert retained.find(qn("w:tblBorders")) is not None
    assert validate_docx_tables(document, profile) == []


@pytest.mark.parametrize("rows", [1, 3])
def test_table_contract_follows_changed_profile_and_turns_off_old_flags(rows):
    document = Document()
    table = document.add_table(rows=rows, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    apply_profile(document)
    profile = load_profile()
    rules = profile["tables"]
    rules["repeat_header_row"] = False
    rules["prevent_row_split"] = False
    rules["inside_horizontal"] = "none"
    for key in ("top_border_ooxml_size", "header_bottom_border_ooxml_size", "bottom_border_ooxml_size"):
        rules[key] += 2
    rules["cell_margins_twips"]["top"] += 13
    apply_profile(document, profile)
    assert validate_docx_tables(document, profile) == []
    assert validate_docx_tables(document)  # The YAML change, not duplicated constants, controls validation.
    assert table.rows[0]._tr.find("./" + qn("w:trPr") + "/" + qn("w:tblHeader")).get(qn("w:val")) == "0"


def test_table_validator_reports_corruption_without_repairing_it():
    document = Document()
    table = document.add_table(rows=3, cols=1)
    apply_profile(document)
    row = table.rows[0]
    row._tr.find("./" + qn("w:trPr") + "/" + qn("w:tblHeader")).set(qn("w:val"), "0")
    _row_border_exception(row)
    properties = table.cell(0, 0)._tc.tcPr
    properties.find("./" + qn("w:tcBorders") + "/" + qn("w:top")).set(qn("w:sz"), "1")
    properties.find("./" + qn("w:tcMar") + "/" + qn("w:start")).set(qn("w:w"), "0")
    before = document.element.xml
    errors = validate_docx_tables(document)
    assert all(any(key in error for error in errors) for key in
               ("tblHeader", "tblPrEx", "border.top", "margin.start"))
    assert document.element.xml == before


def test_apply_profile_itself_rejects_an_incomplete_table_formatter(monkeypatch):
    document = Document()
    document.add_table(rows=2, cols=1)
    monkeypatch.setattr(paper_format, "_format_tables", lambda document, profile: None)
    with pytest.raises(ValueError, match="DOCX profile application failed.*表 1"):
        apply_profile(document)


def test_math_font_and_display_alignment_are_checked_from_profile():
    document = Document()
    paragraph = document.add_paragraph()
    block = OxmlElement("m:oMathPara")
    expression = OxmlElement("m:oMath")
    run = OxmlElement("m:r")
    text = OxmlElement("m:t")
    text.text = "x"
    run.append(text)
    expression.append(run)
    block.append(expression)
    paragraph._p.append(block)
    profile = load_profile()
    profile["equations"].update(math_font="STIX Two Math", display_alignment="right")
    apply_profile(document, profile)
    assert paragraph.alignment == WD_ALIGN_PARAGRAPH.RIGHT
    assert validate_docx_math(document, profile) == []
    assert validate_docx_math(document)
    run.find("./" + qn("w:rPr") + "/" + qn("w:rFonts")).set(qn("w:ascii"), "Segoe Print")
    block.find("./" + qn("m:oMathParaPr") + "/" + qn("m:jc")).set(qn("m:val"), "left")
    errors = validate_docx_math(document, profile)
    assert any("run 1" in error for error in errors)
    assert any("display_alignment" in error for error in errors)


def test_page_grid_distances_and_first_page_follow_profile():
    document = Document()
    profile = load_profile()
    profile["page"]["header_distance_twips"] += 100
    profile["page"]["footer_distance_twips"] += 100
    profile["page"]["document_grid"]["line_pitch_twips"] += 20
    profile["page"]["different_first_page"] = False
    apply_profile(document, profile)
    assert validate_docx_layout(document, profile) == []
    errors = validate_docx_layout(document)
    assert all(any(key in error for error in errors) for key in
               ("header_distance", "footer_distance", "document_grid", "different_first_page"))


def test_saved_docx_is_validated_before_replacing_output(tmp_path, monkeypatch):
    source, target = tmp_path / "source.docx", tmp_path / "candidate.docx"
    document = Document()
    document.add_paragraph("正文")
    document.save(source)
    document.save(target)
    source_bytes = source.read_bytes()
    original_bytes = target.read_bytes()
    save = DocumentClass.save

    def corrupt_serialized_document(self, path):
        save(self, path)
        damaged = Document(path)
        damaged.sections[0].footer_distance += Twips(1)
        save(damaged, path)

    monkeypatch.setattr(DocumentClass, "save", corrupt_serialized_document)
    with pytest.raises(ValueError, match="footer_distance"):
        normalize_docx_file(source, target, overwrite=True)
    assert target.read_bytes() == original_bytes
    assert source.read_bytes() == source_bytes
    assert sorted(path.name for path in tmp_path.iterdir()) == ["candidate.docx", "source.docx"]


def _caption_document():
    document = Document()
    document.add_paragraph("题名", "Title")
    body = document.add_paragraph("中文 English 正文", "Body Text")
    caption = document.add_paragraph("表1 结果", "Caption")
    table = document.add_table(rows=2, cols=1)
    table.cell(0, 0).text = "表头"
    table.cell(1, 0).text = "数据"
    apply_profile(document)
    return document, body, caption, table


@pytest.mark.parametrize("target,attribute,value,expected", [
    ("body", "size", Pt(20), "size_pt"),
    ("body", "first_line_indent", Pt(0), "first_line_indent"),
    ("body", "alignment", WD_ALIGN_PARAGRAPH.LEFT, "alignment"),
    ("body", "space_before", Pt(5), "space_before"),
    ("body", "space_after", Pt(5), "space_after"),
    ("caption", "bold", False, "bold"),
    ("caption", "line_spacing", 2.0, "line_spacing"),
    ("table", "size", Pt(20), "size_pt"),
    ("table", "line_spacing", 2.0, "line_spacing"),
])
def test_six_format_rules_detect_saved_corruption(tmp_path, target, attribute, value, expected):
    document, body, caption, table = _caption_document()
    paragraph = {"body": body, "caption": caption, "table": table.cell(1, 0).paragraphs[0]}[target]
    owner = paragraph.runs[0].font if attribute in {"size", "bold"} else paragraph.paragraph_format
    setattr(owner, attribute, value)
    output = tmp_path / "corrupted.docx"
    document.save(output)
    saved = Document(output)
    errors = validate_docx_layout(saved) + validate_docx_tables(saved)
    assert any(expected in error for error in errors), errors


def test_font_and_character_indent_corruption_are_detected():
    document, body, _, _ = _caption_document()
    body.runs[0]._r.rPr.rFonts.set(qn("w:eastAsia"), "黑体")
    body._p.pPr.ind.set(qn("w:firstLineChars"), "100")
    errors = validate_docx_layout(document)
    assert any("font.eastAsia" in error for error in errors)
    assert any("first_line_indent_chars" in error for error in errors)


def test_inherited_body_format_is_accepted_and_changed_styles_are_detected():
    document, body, _, _ = _caption_document()
    body.runs[0]._r.remove(body.runs[0]._r.rPr)
    assert validate_docx_layout(document) == []
    document.styles["Body Text"].font.size = Pt(20)
    assert any("size_pt" in e for e in validate_docx_layout(document))


def test_caption_position_is_checked_without_moving_content():
    document, _, caption, table = _caption_document()
    table._tbl.addnext(caption._p)
    before = document.element.xml
    errors = paper_format.validate_docx_captions(document)
    assert any("表格上方" in error for error in errors)
    assert document.element.xml == before


def test_caption_duplicates_and_orphans_are_reported():
    document, _, _, _ = _caption_document()
    document.add_paragraph("表1　重复孤立题注", "Table Caption")
    errors = paper_format.validate_docx_captions(document)
    assert any("编号重复" in error for error in errors)
    assert any("表格上方" in error for error in errors)


def test_explicit_caption_with_missing_label_is_not_silently_treated_as_body():
    document = Document()
    document.add_paragraph("缺少图表编号", "Caption")
    with pytest.raises(ValueError, match="缺少可识别的图表编号"):
        apply_profile(document)


def test_figure_caption_must_follow_image_even_when_style_is_explicit():
    document = Document()
    document.add_paragraph("题名", "Title")
    picture = document.add_paragraph()
    picture.add_run()._r.append(OxmlElement("w:drawing"))
    caption = document.add_paragraph("图1 示意", "Caption")
    apply_profile(document)
    assert paper_format.validate_docx_captions(document) == []
    picture._p.addprevious(caption._p)
    assert any("图片下方" in e for e in paper_format.validate_docx_captions(document))


def test_body_references_are_not_captions_even_next_to_table():
    document = Document()
    reference = document.add_paragraph("表1 显示各方案的差异。", "Body Text")
    document.add_table(rows=1, cols=1)
    separate = document.add_paragraph("图2 显示另一趋势。")
    apply_profile(document)
    assert reference.style.name == "Body Text"
    assert separate.style.name == "Normal"
    assert reference.text == "表1 显示各方案的差异。"
    assert reference.alignment == WD_ALIGN_PARAGRAPH.JUSTIFY
    assert paper_format.validate_docx_captions(document) == []


def test_known_pandoc_body_caption_is_marked_from_object_context():
    document = Document()
    caption = document.add_paragraph("表1 结果", "Body Text")
    document.add_table(rows=1, cols=1)
    paper_format.mark_pandoc_captions(document)
    apply_profile(document)
    assert caption.style.name == "Table Caption"
    assert caption.text == "表1" + load_profile()["typography"]["caption"]["label_separator"] + "结果"


def test_corrupt_heading_size_and_caption_spacing_follow_profile():
    document, _, caption, _ = _caption_document()
    heading = document.add_paragraph("1.1 方法", "Heading 2")
    profile = load_profile()
    profile["typography"]["heading_level_2_and_3"]["size_pt"] += 1
    profile["typography"]["caption"]["label_separator"] = "  "
    apply_profile(document, profile)
    assert validate_docx_layout(document, profile) == []
    assert caption.text.startswith("表1  ")
    heading.runs[0].font.size = Pt(10)
    assert any("size_pt" in error for error in validate_docx_layout(document, profile))
