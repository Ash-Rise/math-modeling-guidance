"""Read-only inventory and exploratory spectrum view of the four official inputs."""
from pathlib import Path
import json
import sys
import numpy as np
import openpyxl
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT.parents[3]))
from shared.figure_style import PALETTES

# Same role mapping across spectrum, model comparison and sensitivity views.
FIGURE_COLORS = PALETTES["categorical"]
INPUT = ROOT.parents[1] / 'problem-statements' / 'attachments' / 'problem-b'

def read_data(i):
    path = INPUT / f'附件{i}.xlsx'
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = book.active
    rows = list(sheet.values)
    a = np.asarray(rows[1:], dtype=float)
    info = dict(file=path.name, sheets=book.sheetnames, header=rows[0],
                worksheet_rows=sheet.max_row, worksheet_columns=sheet.max_column,
                records=len(a), first=rows[1], last=rows[-1],
                min=a.min(axis=0).tolist(), max=a.max(axis=0).tolist(),
                missing=int(np.isnan(a).sum()),
                step_range=[float(np.diff(a[:,0]).min()),float(np.diff(a[:,0]).max())])
    book.close()
    return a, info

if __name__ == '__main__':
    import argparse,tempfile
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output-dir',type=Path,help='inventory/preview output; defaults to a new system temporary directory')
    args=parser.parse_args()
    output=args.output_dir or Path(tempfile.mkdtemp(prefix='cumcm-b-inventory-'))
    output=output.resolve()
    output.mkdir(parents=True,exist_ok=True)
    from shared.figure_style import paper_style,figure_size
    with paper_style():
        fig, axs = plt.subplots(2, 1, figsize=figure_size(height=4.5),sharex=True)
        info = []
        for i in range(1,5):
            a, record = read_data(i)
            info.append(record)
            j=(i-1)%2
            axs[(i-1)//2].plot(a[:,0],a[:,1],lw=.95,color=FIGURE_COLORS[j],
                              ls=['-','--'][j],label=f'Attachment {i}')
        for ax, name in zip(axs,['SiC','Si']):
            ax.set(xlabel='Wavenumber (cm$^{-1}$)',ylabel='Reflectance (%)',title=name)
            ax.legend()
        (output/'data-inventory.json').write_text(json.dumps(info,ensure_ascii=False,indent=2),encoding='utf-8')
        fig.tight_layout()
        fig.savefig(output/'spectra-preview.png')
        plt.close(fig)
        print(json.dumps(info,ensure_ascii=False,indent=2))
        print(output)
